import os
import fitz
import docx
import openpyxl
import pytesseract
from PIL import Image

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'


def extract_text_from_pdf(file_path: str) -> str:
    doc = fitz.open(file_path)
    text = ''
    for page in doc:
        text += page.get_text()
    doc.close()
    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    return '\n'.join([para.text for para in doc.paragraphs if para.text.strip()])


def extract_text_from_xlsx(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f'Sheet: {sheet.title}')
        for row in sheet.iter_rows(values_only=True):
            row_text = ' | '.join([str(cell) for cell in row if cell is not None])
            if row_text.strip():
                lines.append(row_text)
    return '\n'.join(lines)


def detect_language_and_extract(file_path: str) -> tuple[str, str]:
    image = Image.open(file_path)

    try:
        osd = pytesseract.image_to_osd(image, output_type=pytesseract.Output.DICT)
        script = osd.get('script', 'Latin')
        script_to_lang = {
            'Cyrillic': 'rus',
            'Latin': 'eng',
            'Arabic': 'ara',
            'Chinese': 'chi_sim',
            'Japanese': 'jpn',
            'Korean': 'kor',
        }
        lang = script_to_lang.get(script, 'eng')
    except Exception:
        lang = 'rus+eng'

    text = pytesseract.image_to_string(image, lang=lang).strip()
    return text, lang


def extract_text_from_image(file_path: str, language: str = 'auto') -> tuple[str, str]:
    if language == 'auto':
        return detect_language_and_extract(file_path)
    lang_map = {'en': 'eng', 'ru': 'rus', 'tj': 'rus'}
    lang = lang_map.get(language, 'eng')
    image = Image.open(file_path)
    text = pytesseract.image_to_string(image, lang=lang).strip()
    return text, lang


def extract_text(file_path: str, file_extension: str) -> str:
    ext = file_extension.lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext in ['.docx', '.doc']:
        return extract_text_from_docx(file_path)
    elif ext in ['.xlsx', '.xls']:
        return extract_text_from_xlsx(file_path)
    elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff']:
        text, _ = extract_text_from_image(file_path)
        return text
    else:
        raise ValueError(f'Unsupported file type: {ext}')


ALLOWED_EXTENSIONS = ['.pdf', '.docx', '.doc', '.xlsx', '.xls', '.jpg', '.jpeg', '.png', '.bmp', '.tiff']
MAX_FILE_SIZE_MB = 20